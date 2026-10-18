"""
Voice Marks → Excel  (Faculty Dictation App)
Works on: Windows 10/11, Mac, iPhone, iPad, Android
"""

import streamlit as st
import pandas as pd
import re
import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
import streamlit.components.v1 as components

st.set_page_config(
    page_title="Voice Marks → Excel",
    page_icon="🎤",
    layout="wide",
    initial_sidebar_state="expanded",
)

if "students" not in st.session_state:
    st.session_state.students = []
if "draft_text" not in st.session_state:
    st.session_state.draft_text = ""


# ------------------------------------------------------------------
# Parser
# ------------------------------------------------------------------
def normalize_transcript(text: str) -> str:
    text = text.lower().replace("\n", " ")
    num_map = {
        "zero": "0", "one": "1", "two": "2", "three": "3", "four": "4",
        "five": "5", "six": "6", "seven": "7", "eight": "8", "nine": "9",
        "ten": "10", "eleven": "11", "twelve": "12", "thirteen": "13",
        "fourteen": "14", "fifteen": "15", "sixteen": "16", "seventeen": "17",
        "eighteen": "18", "nineteen": "19", "twenty": "20",
    }
    for w, d in num_map.items():
        text = re.sub(rf"\b{w}\b", d, text)

    text = text.replace("roll no.", "roll number")
    text = text.replace("roll no", "roll number")
    text = text.replace("ques ", "question ")
    text = text.replace("q.", "question ")
    text = text.replace("question no.", "question ")
    text = text.replace("question number", "question ")
    text = text.replace("question num", "question ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def extract_roll(text: str) -> str:
    patterns = [
        r"\b(2[kK]\d{2}[a-zA-Z]{2,6}\d{3,6})\b",
        r"roll\s*(?:number|no\.?)?\s*[:\-]?\s*([0-9a-zA-Z]{6,16})",
        r"\b(\d{2,4}[a-zA-Z]{2,6}\d{3,6})\b",
    ]
    for pat in patterns:
        m = re.search(pat, text, re.I)
        if m:
            return re.sub(r"\s+", "", m.group(1)).upper()
    return ""


def extract_name(text: str) -> str:
    patterns = [
        r"(?:student\s*)?name\s*(?:is|:)?\s*([a-zA-Z][a-zA-Z\s]{1,30}?)(?:\s*,|\s+roll|\s+question|\s+$)",
        r"(?:name)\s+([a-zA-Z][a-zA-Z\s]{1,25}?)(?:\s+roll)",
    ]
    for pat in patterns:
        m = re.search(pat, text, re.I)
        if m:
            name = re.sub(r"[^a-zA-Z\s]", "", m.group(1)).strip()
            name = re.sub(r"\s+", " ", name)
            if len(name) >= 2 and name.lower() not in ("is", "the", "student"):
                return name.title()
    return ""


def extract_marks(text: str) -> dict:
    marks = {}
    t = text.lower()

    # question 1 a 5 / question 1 (a) 5 / question 1a 5
    for m in re.finditer(
        r"question\s*(\d{1,2})\s*[\(\[\-\s]*([a-j])\s*[\)\]]*\s*[:\-]?\s*(\d{1,2})\b",
        t,
    ):
        marks[f"q{int(m.group(1))}{m.group(2)}"] = int(m.group(3))

    # q1a 5 / q1-a 5
    for m in re.finditer(
        r"\bq\s*(\d{1,2})\s*[\-\s\(\[]*([a-j])\s*[\)\]]*\s*[:\-]?\s*(\d{1,2})\b",
        t,
    ):
        key = f"q{int(m.group(1))}{m.group(2)}"
        if key not in marks:
            marks[key] = int(m.group(3))

    # 1 a 5 marks / 1.(a) 5
    for m in re.finditer(
        r"\b(\d{1,2})\s*[\.\)]\s*[\(\[]?\s*([a-j])\s*[\)\]]?\s*[:\-]?\s*(\d{1,2})\b",
        t,
    ):
        key = f"q{int(m.group(1))}{m.group(2)}"
        if key not in marks:
            marks[key] = int(m.group(3))

    # "1b 3" style after "question number 1"
    for m in re.finditer(r"\b(\d{1,2})\s*([a-j])\s+(\d{1,2})\b", t):
        key = f"q{int(m.group(1))}{m.group(2)}"
        if key not in marks:
            marks[key] = int(m.group(3))

    # question 2 0 / q5 4 (no part)
    for m in re.finditer(r"question\s*(\d{1,2})\s*[:\-]?\s*(\d{1,2})\b", t):
        q = int(m.group(1))
        if not any(k.startswith(f"q{q}") and len(k) > len(f"q{q}") for k in marks):
            marks[f"q{q}"] = int(m.group(2))

    for m in re.finditer(r"\bq\s*(\d{1,2})\s*[:\-]?\s*(\d{1,2})\b", t):
        q = int(m.group(1))
        if f"q{q}" not in marks and not any(
            k.startswith(f"q{q}") and len(k) > len(f"q{q}") for k in marks
        ):
            marks[f"q{q}"] = int(m.group(2))

    return marks


def parse_transcript(raw: str) -> dict:
    text = normalize_transcript(raw)
    name = extract_name(text)
    roll = extract_roll(text)
    marks = extract_marks(text)

    student = {"name": name, "roll": roll, "raw": raw.strip()}
    for q in range(1, 11):
        student[f"q{q}"] = marks.get(f"q{q}")
        for p in "abcdefghij":
            student[f"q{q}{p}"] = marks.get(f"q{q}{p}")

    grand = 0
    for q in range(1, 11):
        q_total = 0
        has_parts = False
        for p in "abcdefghij":
            v = student.get(f"q{q}{p}")
            if v is not None:
                q_total += int(v)
                has_parts = True
        if not has_parts and student.get(f"q{q}") is not None:
            q_total = int(student[f"q{q}"])
        student[f"q{q}_total"] = q_total
        grand += q_total
    student["grand_total"] = grand
    return student


# ------------------------------------------------------------------
# Excel
# ------------------------------------------------------------------
def create_excel(students: list) -> bytes:
    wb = Workbook()
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    green_fill = PatternFill("solid", fgColor="D5F5E3")
    grand_fill = PatternFill("solid", fgColor="F5B7B1")
    yellow_fill = PatternFill("solid", fgColor="FCF3CF")
    thin = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    center = Alignment(horizontal="center", vertical="center")

    ws = wb.active
    ws.title = "Summary"
    for col, h in enumerate(["S.No", "Student Name", "Roll Number", "Grand Total"], 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin
    for i, s in enumerate(students, 1):
        r = i + 1
        ws.cell(row=r, column=1, value=i).border = thin
        ws.cell(row=r, column=2, value=s.get("name", "")).border = thin
        ws.cell(row=r, column=3, value=s.get("roll", "")).border = thin
        cell = ws.cell(row=r, column=4, value=s.get("grand_total", 0))
        cell.fill = grand_fill
        cell.border = thin
        cell.alignment = center
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 12

    ws2 = wb.create_sheet("Student Marks (Flat)")
    flat_headers = ["S.No", "Student Name", "Roll Number"]
    for q in range(1, 11):
        for p in "abcdefghij":
            flat_headers.append(f"{q}({p})")
        flat_headers.append(f"Q{q} Total")
    flat_headers.append("Grand Total")

    for col, h in enumerate(flat_headers, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin

    for i, s in enumerate(students, 1):
        r = i + 1
        col = 1
        ws2.cell(row=r, column=col, value=i).border = thin
        col += 1
        ws2.cell(row=r, column=col, value=s.get("name", "")).border = thin
        col += 1
        ws2.cell(row=r, column=col, value=s.get("roll", "")).border = thin
        col += 1
        for q in range(1, 11):
            part_start = col
            for p in "abcdefghij":
                val = s.get(f"q{q}{p}")
                cell = ws2.cell(row=r, column=col, value=val if val is not None else "")
                cell.border = thin
                cell.alignment = center
                if val is not None:
                    cell.fill = yellow_fill
                col += 1
            q_only = s.get(f"q{q}")
            has_any = any(s.get(f"q{q}{p}") is not None for p in "abcdefghij")
            if not has_any and q_only is not None:
                cell = ws2.cell(row=r, column=col, value=int(q_only))
            else:
                cell = ws2.cell(
                    row=r,
                    column=col,
                    value=f"=SUM({get_column_letter(part_start)}{r}:{get_column_letter(col-1)}{r})",
                )
            cell.fill = green_fill
            cell.border = thin
            cell.alignment = center
            col += 1
        total_cols = []
        c = 4
        for q in range(1, 11):
            total_cols.append(get_column_letter(c + 10))
            c += 11
        cell = ws2.cell(row=r, column=col, value="=" + "+".join(f"{tc}{r}" for tc in total_cols))
        cell.fill = grand_fill
        cell.border = thin
        cell.alignment = center
        cell.font = Font(bold=True)

    ws2.column_dimensions["A"].width = 6
    ws2.column_dimensions["B"].width = 18
    ws2.column_dimensions["C"].width = 16
    for col in range(4, len(flat_headers) + 1):
        ws2.column_dimensions[get_column_letter(col)].width = 7
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("Info")
    ws3["A1"] = "Voice Marks → Excel"
    ws3["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws3["A3"] = f"Students: {len(students)}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ------------------------------------------------------------------
# Sidebar
# ------------------------------------------------------------------
with st.sidebar:
    st.title("🎤 Voice Marks → Excel")
    st.markdown("---")
    st.markdown("### Windows 10 / 11 (laptop)")
    st.markdown(
        """
1. Click **inside the white text box** (important)
2. Press **Windows key + H**
3. If asked, allow microphone
4. Speak clearly
5. Click **Parse & Add Student**

If Win+H does nothing:  
**Settings → Privacy → Speech** → turn on  
**Settings → Time & Language → Speech** → install English
        """
    )
    st.markdown("### Android phone")
    st.markdown(
        """
1. Tap **inside the white text box**
2. Tap **mic on Gboard** (keyboard)
3. Speak → **Parse & Add Student**

Do **not** rely only on “Browser mic” — use Gboard mic in the box.
        """
    )
    st.markdown("### iPhone")
    st.markdown("Tap box → **mic on iOS keyboard** → speak → Parse")
    st.markdown("### Mac")
    st.markdown("Click box → **Fn Fn** → speak → Parse")
    st.markdown("---")
    st.code(
        "Name Rahul Sharma\n"
        "Roll number 2K23CSUN001\n"
        "Question 1 A 5 marks\n"
        "Question 2 zero\n"
        "Question 5 4 marks",
        language=None,
    )


# ------------------------------------------------------------------
# Main
# ------------------------------------------------------------------
st.title("🎤 Voice Marks → Excel")
st.caption("Dictate marks on Windows / Android / iPhone / Mac → Excel with auto-sum")

EXAMPLE = (
    "Name Pratham, roll number 2K23UON0104, "
    "question 1 A 4 marks, question 1 B 4 marks, question 1 C 2 marks, "
    "question 2 zero, question 4 zero, question 5 4 marks, "
    "question 7 2 marks, question 10 4 marks"
)

st.subheader("Step 1 — Put marks text in the box below")

st.info(
    "**Windows:** click the box → press **Win + H** → speak.  \n"
    "**Android:** tap the box → tap **Gboard mic** → speak.  \n"
    "Grey example text is only a hint — your spoken/typed text must appear in the box before Parse."
)

# Single source of truth for text
transcript = st.text_area(
    "Marks script (click here first, then dictate or type)",
    value=st.session_state.draft_text,
    height=150,
    placeholder=EXAMPLE,
    key="marks_input",
)

# Keep draft in sync when user types
if transcript != st.session_state.draft_text:
    st.session_state.draft_text = transcript

c1, c2, c3 = st.columns(3)
with c1:
    parse_btn = st.button("🔍 Parse & Add Student", type="primary", use_container_width=True)
with c2:
    if st.button("📋 Load example", use_container_width=True):
        st.session_state.draft_text = EXAMPLE
        st.rerun()
with c3:
    if st.button("🗑️ Clear text", use_container_width=True):
        st.session_state.draft_text = ""
        st.rerun()

# Optional browser mic — paste helper only
with st.expander("Optional: Browser mic (then paste into the box above)"):
    st.caption(
        "Use this only if Win+H / Gboard is not available. "
        "After listening, Select All → Copy the heard text → paste into the box above → Parse."
    )
    components.html(
        """
<div style="font-family: sans-serif; padding: 8px;">
  <button id="btnStart" style="padding: 8px 14px; background: #1F4E79; color: #fff; border: none; border-radius: 6px; cursor: pointer;">Start browser mic</button>
  <button id="btnStop" style="padding: 8px 14px; margin-left: 6px; border-radius: 6px; cursor: pointer;">Stop</button>
  <span id="status" style="margin-left: 8px; color: #666;"></span>
  <div id="out" contenteditable="true" style="margin-top: 8px; min-height: 50px; padding: 8px; border: 1px solid #ccc; border-radius: 6px; background: #fff;"></div>
</div>
<script>
const out = document.getElementById("out");
const status = document.getElementById("status");
const SR = window.SpeechRecognition || window.webkitSpeechRecognition;
if (!SR) {
  status.textContent = "Not supported — use Win+H or phone keyboard mic in the box above.";
} else {
  const rec = new SR();
  rec.continuous = true;
  rec.interimResults = true;
  rec.lang = "en-IN";
  rec.onresult = (e) => {
    let t = "";
    for (let i = 0; i < e.results.length; i++) t += e.results[i][0].transcript + " ";
    out.innerText = t.trim();
  };
  rec.onerror = (e) => { status.textContent = e.error; };
  rec.onend = () => { status.textContent = "Stopped — copy text into the box above"; };
  document.getElementById("btnStart").onclick = () => {
    out.innerText = "";
    status.textContent = "Listening…";
    try { rec.start(); } catch (err) { status.textContent = String(err); }
  };
  document.getElementById("btnStop").onclick = () => { try { rec.stop(); } catch (e) {} };
}
</script>
        """,
        height=140,
    )

# Parse
if parse_btn:
    text_to_parse = (transcript or st.session_state.draft_text or "").strip()
    if not text_to_parse:
        st.error(
            "The text box is empty. "
            "Click inside the box, use Win+H (Windows) or Gboard mic (Android), "
            "speak, and wait until you see your words in the box — then Parse."
        )
    else:
        student = parse_transcript(text_to_parse)
        if not student["name"] and not student["roll"] and student["grand_total"] == 0:
            st.error(
                "Could not find name, roll, or marks. Try saying:\n\n"
                "Name Rahul Sharma roll number 2K23CSUN001 "
                "question 1 A 5 marks question 2 zero question 5 4 marks"
            )
            with st.expander("What the app heard / read"):
                st.code(text_to_parse)
        else:
            roll = student["roll"]
            if roll:
                st.session_state.students = [
                    s for s in st.session_state.students if s.get("roll") != roll
                ]
            st.session_state.students.append(student)
            st.session_state.draft_text = ""
            st.success(
                f"Added **{student['name'] or 'Unknown'}** "
                f"({student['roll'] or 'no roll'}) — Total: **{student['grand_total']}**"
            )
            st.rerun()

# Live preview when there is text
preview_src = (transcript or st.session_state.draft_text or "").strip()
if preview_src:
    prev = parse_transcript(preview_src)
    with st.expander("Live preview (what will be saved on Parse)", expanded=True):
        a, b, c = st.columns(3)
        a.metric("Name", prev["name"] or "—")
        b.metric("Roll", prev["roll"] or "—")
        c.metric("Grand Total", prev["grand_total"])
        bits = []
        for q in range(1, 11):
            parts = [f"{p}={prev[f'q{q}{p}']}" for p in "abcdefghij" if prev.get(f"q{q}{p}") is not None]
            if parts:
                bits.append(f"Q{q}: {', '.join(parts)}")
            elif prev.get(f"q{q}") is not None:
                bits.append(f"Q{q}: {prev[f'q{q}']}")
        st.write(" · ".join(bits) if bits else "_No marks detected yet_")

# Students table
st.markdown("---")
st.subheader("Step 2 — Students added")

if st.session_state.students:
    rows = []
    for s in st.session_state.students:
        row = {"Name": s.get("name", ""), "Roll No": s.get("roll", ""), "Grand Total": s.get("grand_total", 0)}
        for q in range(1, 11):
            parts = [f"{p}:{s[f'q{q}{p}']}" for p in "abcdefghij" if s.get(f"q{q}{p}") is not None]
            row[f"Q{q}"] = " ".join(parts) if parts else (str(s[f"q{q}"]) if s.get(f"q{q}") is not None else "")
        rows.append(row)
    st.dataframe(pd.DataFrame(rows), use_container_width=True, hide_index=True)

    totals = [s["grand_total"] for s in st.session_state.students]
    m1, m2, m3 = st.columns(3)
    m1.metric("Students", len(totals))
    m2.metric("Highest", max(totals))
    m3.metric("Average", round(sum(totals) / len(totals), 1))

    st.download_button(
        "📥 Download Excel",
        data=create_excel(st.session_state.students),
        file_name=f"Voice_Marks_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
        use_container_width=True,
    )
    if st.button("🔄 Clear all students"):
        st.session_state.students = []
        st.rerun()
else:
    st.info("No students yet. Put text in the box above and click **Parse & Add Student**.")

st.caption(
    "Windows: Win+H inside the box · Android: Gboard mic inside the box · "
    "You must see your words in the white box before Parse will work."
)
